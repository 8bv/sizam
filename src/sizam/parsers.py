import logging
from typing import List, Optional, Iterator

from openpyxl import load_workbook

from .db.models.wiki2035 import UnitWithCourse

logger = logging.getLogger(__name__)


class InvalidHeader(Exception):
    "Неверный формат первой строки в файле"


def get_data_from_excel(
    path: str,
    header_template: List[str],
    sheet: Optional[str] = None,
) -> Iterator:
    workbook = load_workbook(path, read_only=True)
    if sheet is None:
        worksheet = workbook.active
    else:
        worksheet = workbook[sheet]

    wsit = worksheet.iter_rows(values_only=True)

    header = next(wsit)
    if not any(header):
        raise InvalidHeader("Первая строка не найдена")

    if [v.lower() for v in header] != header_template:
        raise InvalidHeader(
            "Формат первой строки не соответствует шаблону: %s"
            % ", ".join(header_template)
        )

    return wsit


def get_units_with_course_from_excel(
    path: str,
    sheet: Optional[str] = None,
) -> List[UnitWithCourse]:
    units_with_course = []
    for row in get_data_from_excel(
        path, header_template=["unti_id", "course_id"], sheet=sheet
    ):
        logger.debug("%s", row)
        if not all(row):
            break

        units_with_course.append(
            UnitWithCourse(unit_id=int(row[0]), course_id=int(row[1]))
        )

    return units_with_course
